"""
extraction/vision_pipeline.py — Multi-format & Vision Document Extraction (Unit 8).

Handles binary document formats (PDF, DOCX, XLSX) and detects all unreadable
failure modes (0-byte, corrupt/garbled, scanned image-only) gracefully without
crashing the pipeline.
"""

import re
import logging
from pathlib import Path
from typing import Optional, Literal
import pypdf
from pypdf.errors import PdfStreamError
import docx
import openpyxl
import json

from shared.schemas import ExtractionResult, ExtractedField
from shared.config import GCP_PROJECT_ID, VERTEX_MODEL
from extraction.text_parser import parse_text_document
from matching.synonym_dict import match_field_label, CANONICAL_FIELDS

logger = logging.getLogger(__name__)

# Tokens that denote missing/omitted required fields
BLANK_TOKENS = {"???", "_______", "TBA", "TBC", "", "N/A", "____MT", "____", "___", "TBD"}

# Wrong document type markers
WRONG_DOC_MARKERS = [
    "COMMERCIAL INVOICE",
    "PACKING LIST",
    "CERTIFICATE OF ORIGIN",
]


def clean_label(lbl: str) -> str:
    """Clean a raw label by stripping CJK / unprintable glyphs and normalizing."""
    if not lbl:
        return ""
    s = re.sub(r"[^\x20-\x7E]+", " ", lbl)
    s = re.sub(r"\s*\(\s*\)", "", s)
    s = re.sub(r"\s*\(\s*", " (", s)
    s = re.sub(r"^TOTAL\s+", "", s, flags=re.I)
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def resolve_field_name(raw_lbl: str) -> Optional[str]:
    """Try matching label with successive normalizations."""
    clbl = clean_label(raw_lbl)
    fname = match_field_label(clbl)
    if fname:
        return fname

    # Try stripping trailing parenthesized segment e.g. (KGS)
    clbl2 = re.sub(r"\s*\([^)]*\)$", "", clbl).strip()
    fname = match_field_label(clbl2)
    if fname:
        return fname

    return match_field_label(raw_lbl)



def is_wrong_doc_type(text: str) -> bool:
    """Check if text contains markers for wrong document types."""
    upper = text.upper()
    return any(marker in upper for marker in WRONG_DOC_MARKERS)


def is_blank_value(val: Optional[str]) -> bool:
    """Check if value is considered missing or blank."""
    if val is None:
        return True
    cleaned = val.strip()
    return cleaned in BLANK_TOKENS or not cleaned


# ── XLSX Parser ───────────────────────────────────────────────────────────

def extract_from_xlsx(file_path: Path, doc_type: str, email_id: str) -> ExtractionResult:
    """Extract fields from an XLSX document using openpyxl."""
    try:
        wb = openpyxl.load_workbook(str(file_path), data_only=True)
        ws = wb.active
    except Exception as e:
        logger.warning("Failed to open XLSX %s: %s", file_path, e)
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="unreadable")

    extracted: dict[str, tuple[str, str]] = {}  # field_name -> (value, snippet)
    full_text_chunks = []

    for row in ws.iter_rows():
        vals = [c.value for c in row]
        if not vals:
            continue
        row_str = " ".join(str(v) for v in vals if v is not None)
        full_text_chunks.append(row_str)

        if len(vals) >= 2 and vals[0] is not None:
            raw_lbl = str(vals[0]).strip()
            val = str(vals[1]).strip() if vals[1] is not None else ""
            fname = resolve_field_name(raw_lbl)
            if fname and fname in CANONICAL_FIELDS and fname not in extracted:
                extracted[fname] = (val, f"{raw_lbl}: {val}")

    full_text = "\n".join(full_text_chunks)
    if is_wrong_doc_type(full_text):
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="wrong_doc_type")

    fields = []
    for fname in CANONICAL_FIELDS:
        if fname in extracted:
            val, snip = extracted[fname]
            if is_blank_value(val):
                fields.append(ExtractedField(field_name=fname, value=None, source_snippet=snip, confidence=0.0))
            else:
                fields.append(ExtractedField(field_name=fname, value=val, source_snippet=snip, confidence=1.0))
        else:
            fields.append(ExtractedField(field_name=fname, value=None, source_snippet=None, confidence=0.0))

    return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=fields, doc_status="ok")


# ── DOCX Parser ───────────────────────────────────────────────────────────

def extract_from_docx(file_path: Path, doc_type: str, email_id: str) -> ExtractionResult:
    """Extract fields from a DOCX document using python-docx."""
    try:
        doc = docx.Document(str(file_path))
    except Exception as e:
        logger.warning("Failed to open DOCX %s: %s", file_path, e)
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="unreadable")

    full_text_chunks = [p.text for p in doc.paragraphs]
    extracted: dict[str, tuple[str, str]] = {}

    for t in doc.tables:
        for r in t.rows:
            if len(r.cells) >= 2:
                raw_lbl = r.cells[0].text.strip()
                val = r.cells[1].text.strip()
                full_text_chunks.append(f"{raw_lbl}: {val}")

                fname = resolve_field_name(raw_lbl)
                if fname and fname in CANONICAL_FIELDS and fname not in extracted:
                    extracted[fname] = (val, f"{raw_lbl}: {val}")

    full_text = "\n".join(full_text_chunks)
    if is_wrong_doc_type(full_text):
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="wrong_doc_type")

    fields = []
    for fname in CANONICAL_FIELDS:
        if fname in extracted:
            val, snip = extracted[fname]
            if is_blank_value(val):
                fields.append(ExtractedField(field_name=fname, value=None, source_snippet=snip, confidence=0.0))
            else:
                fields.append(ExtractedField(field_name=fname, value=val, source_snippet=snip, confidence=1.0))
        else:
            fields.append(ExtractedField(field_name=fname, value=None, source_snippet=None, confidence=0.0))

    return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=fields, doc_status="ok")


# ── PDF Parser ────────────────────────────────────────────────────────────

def parse_pdf_text_stream(text: str, doc_type: str, email_id: str) -> Optional[ExtractionResult]:
    """Parse text extracted from digital PDF using block structure."""
    if is_wrong_doc_type(text):
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="wrong_doc_type")

    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    extracted: dict[str, tuple[str, str]] = {}
    current_field = None
    current_vals = []

    for line in lines:
        if ":" in line:
            parts = line.split(":", 1)
            raw_lbl = parts[0].strip()
            fname = resolve_field_name(raw_lbl)
            if fname and fname in CANONICAL_FIELDS:
                val = parts[1].strip()
                extracted[fname] = (val, line)
                current_field = None
                current_vals = []
                continue

        fname = resolve_field_name(line)
        if fname and fname in CANONICAL_FIELDS:
            if current_field and current_vals:
                combined_val = "\n".join(current_vals)
                extracted[current_field] = (combined_val, f"{current_field}: {combined_val}")
            current_field = fname
            current_vals = []
        elif current_field:
            if line.startswith("CONTAINER NO") or line.startswith("HS CODE") or line.startswith("B/L NUMBER"):
                combined_val = "\n".join(current_vals)
                extracted[current_field] = (combined_val, f"{current_field}: {combined_val}")
                current_field = None
                current_vals = []
            else:
                current_vals.append(line)

    if current_field and current_vals:
        combined_val = "\n".join(current_vals)
        extracted[current_field] = (combined_val, f"{current_field}: {combined_val}")

    # Build ExtractedFields list
    fields = []
    for fname in CANONICAL_FIELDS:
        if fname in extracted:
            val, snip = extracted[fname]
            if is_blank_value(val):
                fields.append(ExtractedField(field_name=fname, value=None, source_snippet=snip, confidence=0.0))
            else:
                fields.append(ExtractedField(field_name=fname, value=val, source_snippet=snip, confidence=1.0))
        else:
            fields.append(ExtractedField(field_name=fname, value=None, source_snippet=None, confidence=0.0))

    return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=fields, doc_status="ok")


def extract_pdf_with_gemini(file_path: Path, doc_type: str, email_id: str) -> ExtractionResult:
    """Extract fields from digital PDF using Gemini 2.5 Flash multimodal on Vertex AI."""
    from google import genai
    from google.genai import types

    client = genai.Client(vertexai=True, project=GCP_PROJECT_ID, location="us-central1")
    pdf_bytes = file_path.read_bytes()

    prompt = """\
Extract the 7 canonical shipping document fields from this document.
Fields to extract:
1. shipper (name and address)
2. consignee (name and address)
3. notify_party (name and address)
4. port_of_loading
5. port_of_discharge
6. container_count (e.g. 6 x 40'HC)
7. gross_weight_kg (total gross weight in KG)

If this document is a Commercial Invoice, Packing List, or Certificate of Origin, set "doc_status": "wrong_doc_type".
Otherwise set "doc_status": "ok".

Return ONLY JSON in this exact shape:
{
  "doc_status": "ok",
  "fields": {
    "shipper": "<value or null>",
    "consignee": "<value or null>",
    "notify_party": "<value or null>",
    "port_of_loading": "<value or null>",
    "port_of_discharge": "<value or null>",
    "container_count": "<value or null>",
    "gross_weight_kg": "<value or null>"
  }
}
"""
    try:
        from escalation.feedback_loop import inject_few_shot_into_extraction_prompt
        prompt = inject_few_shot_into_extraction_prompt(prompt)
    except Exception as e:
        logger.debug("Could not inject few-shot guidance: %s", e)

    response = client.models.generate_content(
        model=VERTEX_MODEL,
        contents=[
            types.Part.from_bytes(data=pdf_bytes, mime_type="application/pdf"),
            prompt,
        ],
    )

    raw_text = response.text.strip()
    if raw_text.startswith("```"):
        raw_text = raw_text.split("\n", 1)[1]
        raw_text = raw_text.rsplit("```", 1)[0]
        raw_text = raw_text.strip()

    data = json.loads(raw_text)
    doc_status = data.get("doc_status", "ok")
    if doc_status == "wrong_doc_type":
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="wrong_doc_type")

    parsed_fields = data.get("fields", {})
    fields = []
    for fname in CANONICAL_FIELDS:
        val = parsed_fields.get(fname)
        if is_blank_value(val):
            fields.append(ExtractedField(field_name=fname, value=None, source_snippet=None, confidence=0.0))
        else:
            fields.append(ExtractedField(field_name=fname, value=str(val), source_snippet=str(val), confidence=0.95))

    return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=fields, doc_status="ok")


def extract_from_pdf(file_path: Path, doc_type: str, email_id: str) -> ExtractionResult:
    """Handle PDF extraction, detecting all 3 unreadable failure modes."""
    # 1. Check 0-byte file
    if file_path.stat().st_size == 0:
        logger.info("PDF %s is 0 bytes -> unreadable", file_path.name)
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="unreadable")

    # 2. Open and inspect with pypdf
    try:
        reader = pypdf.PdfReader(str(file_path))
        raw_text = "".join(page.extract_text() or "" for page in reader.pages).strip()
    except (PdfStreamError, Exception) as e:
        logger.info("PDF %s is corrupt/truncated (%s) -> unreadable", file_path.name, type(e).__name__)
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="unreadable")

    # 3. Check for image-only scanned PDF (no extractable text layer)
    if not raw_text:
        logger.info("PDF %s has no text layer (scanned copy) -> unreadable", file_path.name)
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="unreadable")

    # 4. Check for wrong document type
    if is_wrong_doc_type(raw_text):
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="wrong_doc_type")

    # 5. Extract digital PDF via block parser
    res = parse_pdf_text_stream(raw_text, doc_type, email_id)
    missing_fields = [f for f in res.fields if f.value is None]
    if len(missing_fields) > 2:
        try:
            return extract_pdf_with_gemini(file_path, doc_type, email_id)
        except Exception as e:
            logger.warning("Gemini PDF fallback failed for %s: %s", file_path.name, e)

    return res


# ── Master Document Dispatcher ────────────────────────────────────────────

def extract_document(
    file_path: Path,
    doc_type: str,
    email_id: str,
) -> ExtractionResult:
    """Master document extraction entry point. Dispatches by file format."""
    if not file_path.exists() or file_path.stat().st_size == 0:
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="unreadable")

    ext = file_path.suffix.lower()

    try:
        if ext == ".txt":
            text = file_path.read_text(encoding="utf-8", errors="replace")
            res = parse_text_document(text, doc_type=doc_type, email_id=email_id)

        elif ext == ".docx":
            res = extract_from_docx(file_path, doc_type=doc_type, email_id=email_id)

        elif ext == ".xlsx":
            res = extract_from_xlsx(file_path, doc_type=doc_type, email_id=email_id)

        elif ext == ".pdf":
            res = extract_from_pdf(file_path, doc_type=doc_type, email_id=email_id)

        else:
            logger.warning("Unsupported file extension %s for %s", ext, file_path.name)
            return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="unreadable")

        # Feedback loop: apply human operator correction overrides if active
        if res.doc_status == "ok" and res.fields:
            try:
                from escalation.feedback_loop import apply_correction_overrides
                res.fields = apply_correction_overrides(res.fields)
            except Exception as e:
                logger.debug("Failed applying correction overrides: %s", e)

        return res

    except Exception as e:
        logger.error("Unhandled exception extracting %s: %s", file_path.name, e)
        return ExtractionResult(email_id=email_id, doc_type=doc_type, fields=[], doc_status="unreadable")
